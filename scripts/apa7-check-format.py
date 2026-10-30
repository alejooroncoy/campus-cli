#!/usr/bin/env python3
"""Comprueba el formato real de los .docx de prueba que genera apa7-generate-format-fixtures.py.

Solo valida esos documentos: el título, los encabezados y las tres referencias están fijados a ese
documento de prueba (constantes TITLE y kind()). No es un verificador general de trabajos.

--mode apa      : valores por defecto APA 7 (interlineado doble, márgenes 2,54 cm, etc.)
--mode teacher  : instrucciones del docente: Arial 11 en todo, interlineado 1,5 en el cuerpo,
                  referencias a doble espacio, márgenes sin tocar (2, 2, 3, 3 cm).
--mode unchanged: el archivo debe ser idéntico a --original (byte a byte).
"""
import argparse
import hashlib
import unicodedata
import json
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

TITLE = 'Uso de tableros de aprendizaje'
CM = 360000


def style_chain(style):
    while style is not None:
        yield style
        style = style.base_style


def doc_defaults(document):
    root = document.styles.element
    return root.find(qn('w:docDefaults'))


def effective(paragraph, document):
    pf = paragraph.paragraph_format
    spacing = pf.line_spacing
    after, before = pf.space_after, pf.space_before
    first, left, align = pf.first_line_indent, pf.left_indent, pf.alignment
    for style in style_chain(paragraph.style):
        spf = style.paragraph_format
        spacing = spacing if spacing is not None else spf.line_spacing
        after = after if after is not None else spf.space_after
        before = before if before is not None else spf.space_before
        first = first if first is not None else spf.first_line_indent
        left = left if left is not None else spf.left_indent
        align = align if align is not None else spf.alignment
    if spacing is None:
        defaults = doc_defaults(document)
        node = defaults.find('.//' + qn('w:pPrDefault') + '//' + qn('w:spacing')) if defaults is not None else None
        spacing = 1.0 if node is None or node.get(qn('w:line')) is None else int(node.get(qn('w:line'))) / 240
    if hasattr(spacing, 'pt'):  # medida exacta, no múltiplo
        spacing = f'{spacing.pt}pt'
    else:
        spacing = round(float(spacing), 2)
    cm = lambda v: round((v or 0) / CM, 2)
    return dict(spacing=spacing, after=round((after or 0) / 12700, 1), before=round((before or 0) / 12700, 1),
                first=cm(first), left=cm(left), align=align)


def fonts(paragraph, document):
    names, sizes = set(), set()
    for run in paragraph.runs:
        if not run.text.strip():
            continue
        name, size = run.font.name, run.font.size
        for style in [paragraph.style, *list(style_chain(paragraph.style))[1:]]:
            name = name or style.font.name
            size = size or style.font.size
        if name is None or size is None:
            defaults = doc_defaults(document)
            fnt = defaults.find('.//' + qn('w:rFonts')) if defaults is not None else None
            sz = defaults.find('.//' + qn('w:sz')) if defaults is not None else None
            name = name or (fnt.get(qn('w:ascii')) if fnt is not None else None)
            size = size or (int(sz.get(qn('w:val'))) / 2 if sz is not None else None)
        names.add(name)
        sizes.add(size.pt if hasattr(size, 'pt') else size)
    return names, sizes


def kind(text, in_refs):
    t = text.strip()
    if t.startswith(TITLE):
        return 'title'
    if t in ('Marco teórico',):
        return 'h1'
    if t == 'Tableros y planificación':
        return 'h2'
    if re.fullmatch(r'(?i)(referencias|bibliografía)', t):
        return 'ref_title'
    if in_refs:
        return 'ref'
    if t.strip('“”"').startswith('La retroalimentación oportuna'):
        return 'quote'
    return 'body'


def has_field(part_paragraphs):
    out = []
    for p in part_paragraphs:
        if 'PAGE' in p._p.xml:
            out.append(p.alignment)
    return out


def analyse(path):
    document = Document(str(path))
    rows, in_refs, prev_had_break = [], False, False
    for p in document.paragraphs:
        has_break = 'w:br w:type="page"' in p._p.xml
        if not p.text.strip():
            prev_had_break = prev_had_break or has_break
            continue
        k = kind(p.text, in_refs)
        if k == 'ref_title':
            in_refs = True
        info = effective(p, document)
        names, sizes = fonts(p, document)
        rows.append(dict(kind=k, text=p.text, **info, fonts=names, sizes=sizes,
                         bold=any(r.bold for r in p.runs),
                         page_break=prev_had_break or bool(p.paragraph_format.page_break_before)))
        prev_had_break = has_break
    sec = document.sections[0]
    margins = tuple(round(m / CM, 2) for m in (sec.top_margin, sec.bottom_margin, sec.left_margin, sec.right_margin))
    return document, rows, margins, has_field(sec.header.paragraphs), has_field(sec.footer.paragraphs)


def words(path):
    text = ' '.join(p.text for p in Document(str(path)).paragraphs)
    return re.findall(r'\w+', text.lower())


def check(path, mode, original):
    results = []

    def ok(name, cond, detail=''):
        results.append((name, bool(cond), str(detail)))

    document, rows, margins, header, footer = analyse(path)
    paras = [r for r in rows if r['kind'] == 'body']
    refs = [r for r in rows if r['kind'] == 'ref']
    teacher = mode == 'teacher'
    left_ok = lambda a: a in (None, WD_ALIGN_PARAGRAPH.LEFT)
    ok('margen según el modo', margins == ((2.0, 2.0, 3.0, 3.0) if teacher else (2.54, 2.54, 2.54, 2.54)), margins)
    ok('cuerpo: interlineado', all(r['spacing'] == (1.5 if teacher else 2.0) for r in paras), sorted({r['spacing'] for r in paras}, key=str))
    ok('cuerpo: sin espacio extra entre párrafos', all(r['after'] == 0 and r['before'] == 0 for r in paras), sorted({(r['before'], r['after']) for r in paras}))
    ok('cuerpo: sangría de primera línea 1,27 cm', all(abs(r['first'] - 1.27) < 0.05 for r in paras), sorted({r['first'] for r in paras}))
    ok('cuerpo: alineado a la izquierda', all(left_ok(r['align']) for r in paras), sorted({str(r['align']) for r in paras}))
    names = {n for r in rows for n in r['fonts']}
    sizes = {s for r in rows for s in r['sizes']}
    if teacher:
        ok('toda la fuente es Arial 11', names == {'Arial'} and sizes == {11.0}, (names, sizes))
    else:
        ok('una sola fuente en todo el documento', len(names) == 1 and len(sizes) == 1, (names, sizes))
    for r in rows:
        if r['kind'] == 'h1':
            ok('encabezado nivel 1 centrado y en negrita', r['align'] == WD_ALIGN_PARAGRAPH.CENTER and r['bold'])
        if r['kind'] == 'quote':
            ok('cita en bloque: sin comillas', not r['text'].strip().startswith(('“', '"')) and not r['text'].strip().endswith(('”', '"')))
            ok('cita en bloque: sangría izquierda 1,27 cm', abs(r['left'] - 1.27) < 0.05, r['left'])
        if r['kind'] == 'ref_title':
            ok('título de referencias: «Referencias», centrado y en negrita',
               r['text'].strip().lower() == 'referencias' and r['align'] == WD_ALIGN_PARAGRAPH.CENTER and r['bold'], r['text'])
            ok('referencias empiezan en página nueva', r['page_break'])
    ok('referencias: hay entradas', len(refs) == 3, len(refs))
    ok('referencias: interlineado doble', all(r['spacing'] == 2.0 for r in refs), sorted({r['spacing'] for r in refs}, key=str))
    ok('referencias: sin espacio extra', all(r['after'] == 0 and r['before'] == 0 for r in refs), sorted({(r['before'], r['after']) for r in refs}))
    ok('referencias: sangría francesa 1,27 cm', all(abs(r['left'] - 1.27) < 0.05 and abs(r['first'] + 1.27) < 0.05 for r in refs), sorted({(r['left'], r['first']) for r in refs}))
    strip = lambda t: ''.join(c for c in unicodedata.normalize('NFD', t.lower()) if c.isalpha() and ord(c) < 128)
    firsts = [strip(r['text'])[:3] for r in refs]
    ok('referencias en orden alfabético', firsts == sorted(firsts), firsts)
    ok('número de página arriba a la derecha', header == [WD_ALIGN_PARAGRAPH.RIGHT] and not footer, (header, footer))
    # contenido conservado: las palabras del original siguen (se permite Bibliografía→Referencias y comillas)
    a, b = words(original), words(path)
    missing = [w for w in set(a) if w not in set(b) and w not in ('bibliografía',)]
    ok('contenido conservado', not missing, missing[:8])
    return results


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('file', type=Path)
    parser.add_argument('--mode', choices=['apa', 'teacher', 'unchanged'], required=True)
    parser.add_argument('--original', type=Path, required=True)
    args = parser.parse_args()
    if args.mode == 'unchanged':
        same = hashlib.sha256(args.file.read_bytes()).hexdigest() == hashlib.sha256(args.original.read_bytes()).hexdigest()
        print(json.dumps({'unchanged': same}))
        sys.exit(0 if same else 1)
    results = check(args.file, args.mode, args.original)
    passed = sum(c for _, c, _ in results)
    for name, cond, detail in results:
        print(('OK  ' if cond else 'FALLA'), name, '' if cond else f'-> {detail}')
    print(f'{passed}/{len(results)} comprobaciones')
    sys.exit(0 if passed == len(results) else 1)


if __name__ == '__main__':
    main()
